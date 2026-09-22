"""Audit the Mi4/C3 scope of the Gür et al. 2024 public data release."""

from __future__ import annotations

import json
import re
import struct
import subprocess
from pathlib import Path

import yaml
from openpyxl import load_workbook

from fly_emotion.driving.v7_geometry_sign import _sha256

CONFIG = Path("configs/driving-v7-stable-contrast-source-scope-audit.yaml")
IMPLEMENTATION = Path("src/fly_emotion/driving/v7_stable_contrast_source_scope_audit.py")


def _verify(root: Path, spec: dict) -> Path:
    path = root / spec["path"]
    if path.stat().st_size != int(spec["bytes"]) or _sha256(path) != spec["sha256"]:
        raise ValueError(f"stable-contrast snapshot identity changed: {path}")
    return path


def _workbook_inventory(path: Path, expected_fields: list[str]) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = [
        tuple(row)
        for row in workbook.active.iter_rows(values_only=True)
        if any(value is not None for value in row)
    ]
    fields = list(rows[0])
    if fields != expected_fields:
        raise ValueError(f"proofreading workbook fields changed: {path}")
    padded = [row + (None,) * (len(fields) - len(row)) for row in rows[1:]]
    field_text = " ".join(str(value) for row in padded for value in row if value is not None)
    return {
        "sheet_names": workbook.sheetnames,
        "fields": fields,
        "row_count": len(padded),
        "hemispheres": sorted({row[1] for row in padded}),
        "unique_segment_ID_count": len({row[2] for row in padded}),
        "unique_optic_lobe_ID_count": len({row[4] for row in padded}),
        "proofreading_statuses": sorted({row[5] for row in padded}),
        "all_fields_complete": all(value is not None for row in padded for value in row),
        "time_field_found": bool(re.search(r"time|second|frame", " ".join(fields), re.I)),
        "response_field_found": bool(
            re.search(r"response|voltage|fluorescence", " ".join(fields), re.I)
        ),
        "stimulus_field_found": bool(
            re.search(r"stimulus|contrast|luminance", " ".join(fields), re.I)
        ),
        "recording_term_found_in_values": bool(
            re.search(r"recording|GCaMP|voltage", field_text, re.I)
        ),
    }


def _central_directory_names(path: Path) -> list[str]:
    payload = path.read_bytes()
    offset = 0
    names = []
    while offset + 46 <= len(payload) and payload[offset : offset + 4] == b"PK":
        values = struct.unpack_from("<4s6H3L5H2L", payload, offset)
        filename_length, extra_length, comment_length = values[10:13]
        start = offset + 46
        names.append(payload[start : start + filename_length].decode("utf-8", errors="replace"))
        offset = start + filename_length + extra_length + comment_length
    if offset != len(payload):
        raise ValueError("processed ZIP central directory parsing incomplete")
    return names


def evaluate_v7_stable_contrast_source_scope_audit(root: Path) -> dict:
    config = yaml.safe_load((root / CONFIG).read_text(encoding="utf-8"))
    paths = {name: _verify(root, spec) for name, spec in config["snapshots"].items()}
    paper = config["paper"]
    crossref = json.loads(paths["crossref"].read_text())["message"]
    if crossref["DOI"] != paper["doi"]:
        raise ValueError("stable-contrast paper identity changed")
    zenodo = json.loads(paths["zenodo"].read_text())
    if zenodo["doi"] != config["zenodo"]["record_doi"]:
        raise ValueError("stable-contrast Zenodo identity changed")
    file_specs = {item["key"]: item for item in zenodo["files"]}
    if set(file_specs) != {"processed_data.zip", "raw_data.zip"}:
        raise ValueError("stable-contrast Zenodo top-level file inventory changed")
    for name, prefix in (("processed_data.zip", "processed_data"), ("raw_data.zip", "raw_data")):
        item = file_specs[name]
        if (
            int(item["size"]) != int(config["zenodo"][f"{prefix}_bytes"])
            or item["checksum"] != f"md5:{config['zenodo'][f'{prefix}_md5']}"
        ):
            raise ValueError(f"stable-contrast Zenodo file identity changed: {name}")

    html = paths["pmc_article"].read_text(encoding="utf-8")
    text = " ".join(re.sub(r"<[^>]+>", " ", html).split())
    required_phrases = (
        "dendrites of two third-order neurons, Tm1 and Tm9",
        "calcium indicator GCaMP6f",
        "Source data of this study can be found on Zenodo",
        "https://github.com/silieslab/Gur-etal-2024",
    )
    if any(phrase not in text for phrase in required_phrases):
        raise ValueError("stable-contrast paper evidence changed")
    source_term_hits = {
        term: len(re.findall(rf"(?<![A-Za-z0-9]){term}(?![A-Za-z0-9])", text, re.I))
        for term in ("Mi4", "C3", "Tm1", "Tm2", "Tm4", "Tm9", "Dm12")
    }
    if source_term_hits["Mi4"] or source_term_hits["C3"]:
        raise ValueError("stable-contrast paper gained Mi4 or C3 text")

    names = _central_directory_names(paths["processed_central_directory"])
    if len(names) != int(config["zenodo"]["processed_zip_entry_count"]):
        raise ValueError("stable-contrast processed ZIP entry count changed")
    source_paths = {
        source: sorted(name for name in names if source in name)
        for source in ("Mi4", "C3")
    }
    workbooks = {
        source: _workbook_inventory(
            paths[f"{source}_proofreadings"], config["proofreading_fields"]
        )
        for source in ("Mi4", "C3")
    }

    repository = root / config["repository"]["path"]
    local_head = subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=repository, check=True,
        capture_output=True, text=True
    ).stdout.strip()
    remote_head = subprocess.run(
        ["git", "rev-parse", "refs/remotes/origin/main"], cwd=repository,
        check=True, capture_output=True, text=True
    ).stdout.strip()
    tree_files = subprocess.run(
        ["git", "ls-tree", "-r", "--name-only", "HEAD"], cwd=repository,
        check=True, capture_output=True, text=True
    ).stdout.splitlines()
    commit_count = int(subprocess.run(
        ["git", "rev-list", "--all", "--count"], cwd=repository, check=True,
        capture_output=True, text=True
    ).stdout.strip())
    repo_spec = config["repository"]
    if (
        local_head != repo_spec["current_remote_head"]
        or remote_head != local_head
        or len(tree_files) != int(repo_spec["tree_file_count"])
        or commit_count != int(repo_spec["commit_count"])
    ):
        raise ValueError("stable-contrast repository identity changed")

    return {
        "protocol": {
            "name": config["name"],
            "observed_on": config["observed_on"],
            "dependencies_sha256": {
                str(CONFIG): _sha256(root / CONFIG),
                str(IMPLEMENTATION): _sha256(root / IMPLEMENTATION),
            },
            "raw_snapshot_identity": {
                str(path.relative_to(root)): _sha256(path) for path in paths.values()
            },
            "parameter_fit": False,
            "target_activity_injection": False,
            "runtime_modified": False,
        },
        "paper": paper,
        "paper_source_term_hits": source_term_hits,
        "directly_recorded_neuron_types": [
            "L1", "L2", "L3", "Tm1", "Tm2", "Tm4", "Tm9", "Dm12", "T4/T5"
        ],
        "measurement_modalities": ["GCaMP6f_calcium", "iGluSnFR_glutamate"],
        "Mi4_direct_physiology_found": False,
        "C3_direct_physiology_found": False,
        "Zenodo": {
            "concept_doi": config["zenodo"]["concept_doi"],
            "record_doi": zenodo["doi"],
            "top_level_files": {
                name: {"bytes": int(item["size"]), "checksum": item["checksum"]}
                for name, item in file_specs.items()
            },
            "processed_ZIP_entry_count": len(names),
            "processed_source_paths": source_paths,
            "full_processed_archive_downloaded": False,
            "full_raw_archive_downloaded": False,
        },
        "proofreading_workbooks": workbooks,
        "Mi4_C3_files_are_connectome_proofreading_only": all(
            not item["time_field_found"]
            and not item["response_field_found"]
            and not item["stimulus_field_found"]
            and not item["recording_term_found_in_values"]
            for item in workbooks.values()
        ),
        "external_recording_to_MaleCNS_crosswalk_found": False,
        "repository": {
            "url": repo_spec["url"],
            "local_and_remote_head": local_head,
            "commit_count": commit_count,
            "tree_file_count": len(tree_files),
        },
        "authorize_Mi4_C3_source_dynamics_transfer": False,
        "authorize_T4_source_dynamics_fit": False,
        "advance_to_T4_functional_precheck": False,
        "advance_to_LPLC_mechanism_repair": False,
        "advance_to_runtime_integration": False,
        "stop_reason": "Mi4_C3_named_files_are_anatomical_proofreading_not_physiology",
        "boundary": config["boundary"],
    }
