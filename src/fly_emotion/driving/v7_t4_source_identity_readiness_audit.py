"""Audit whether T4 source-voltage evidence supports an individual-level split."""

from __future__ import annotations

import json
import re
from pathlib import Path

import yaml
from openpyxl import load_workbook

from fly_emotion.driving.v7_geometry_sign import _sha256

CONFIG = Path("configs/driving-v7-t4-source-identity-readiness-audit.yaml")
IMPLEMENTATION = Path("src/fly_emotion/driving/v7_t4_source_identity_readiness_audit.py")


def evaluate_v7_t4_source_identity_readiness_audit(root: Path) -> dict:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "the T4 identity audit requires the pypdf development dependency"
        ) from exc
    config = yaml.safe_load((root / CONFIG).read_text(encoding="utf-8"))
    workbook_spec = config["source_workbook"]
    workbook_path = root / workbook_spec["path"]
    if workbook_path.stat().st_size != int(workbook_spec["bytes"]):
        raise ValueError("T4 source workbook size mismatch")
    if _sha256(workbook_path) != workbook_spec["sha256"]:
        raise ValueError("T4 source workbook SHA-256 mismatch")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if workbook.sheetnames != workbook_spec["expected_sheets"]:
        raise ValueError("T4 source workbook sheets changed")
    if any(workbook[name].sheet_state != "visible" for name in workbook.sheetnames):
        raise ValueError("T4 source workbook contains a hidden sheet")
    paper_spec = config["paper_pdf"]
    paper_path = root / paper_spec["path"]
    if paper_path.stat().st_size != int(paper_spec["bytes"]):
        raise ValueError("T4 paper PDF size mismatch")
    if _sha256(paper_path) != paper_spec["sha256"]:
        raise ValueError("T4 paper PDF SHA-256 mismatch")
    paper = PdfReader(paper_path)
    if len(paper.pages) != int(paper_spec["pages"]):
        raise ValueError("T4 paper PDF page count changed")
    paper_text = " ".join(
        " ".join((page.extract_text() or "").replace("- ", "").split()) for page in paper.pages
    )
    one_cell_per_animal = all(
        phrase in paper_text
        for phrase in (
            "Sample sizes in electrophysiological experiments",
            "number of cells, each of which was recorded in a different animal",
        )
    )
    if not one_cell_per_animal:
        raise ValueError("T4 paper no longer states one electrophysiology cell per animal")
    identity_pattern = re.compile(
        "|".join(re.escape(token) for token in config["identity_tokens"]), re.I
    )
    identity_fields = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and identity_pattern.search(cell.value):
                    identity_fields.append(
                        {"sheet": sheet_name, "cell": cell.coordinate, "value": cell.value}
                    )
    source_headers = {}
    for sheet_name in config["source_header_sheets"]:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        source_headers[sheet_name] = {}
        for source_type, count in config["source_types"].items():
            actual = [value for value in headers if str(value).startswith(f"{source_type}-")]
            expected = [f"{source_type}-{index}" for index in range(1, int(count) + 1)]
            if actual != expected:
                raise ValueError(f"T4 anonymous cell ordinals changed for {source_type}")
            source_headers[sheet_name][source_type] = actual
    on_headers = source_headers[config["source_header_sheets"][0]]
    off_headers = source_headers[config["source_header_sheets"][1]]
    paired_ordinals = {source: on_headers[source] == off_headers[source] for source in on_headers}
    if not all(paired_ordinals.values()):
        raise ValueError("T4 ON/OFF cell ordinals are not paired")
    paths = {
        name: Path(config[name])
        for name in (
            "ephys_evidence",
            "kernel_evidence",
            "robustness_evidence",
            "external_contract",
        )
    }
    reports = {
        name: json.loads((root / path).read_text(encoding="utf-8")) for name, path in paths.items()
    }
    ephys = reports["ephys_evidence"]
    verified_files = sorted(ephys["verified_files"])
    if verified_files != sorted(config["expected_verified_edmond_files"]):
        raise ValueError("verified Edmond subset changed")
    identity_sidecars = [
        name
        for name in verified_files
        if identity_pattern.search(name)
        or any(token in name.lower() for token in ("meta", "manifest", "index"))
    ]
    kernel = reports["kernel_evidence"]
    robustness = reports["robustness_evidence"]
    kernel_ready = {
        condition: {
            source: bool(item["ready"])
            for source, item in details["source_results"].items()
            if source in config["source_types"]
        }
        for condition, details in kernel["conditions"].items()
    }
    robustness_ready = {
        condition: {
            source: bool(item["passed"])
            for source, item in details["sources"].items()
            if source in config["source_types"]
        }
        for condition, details in robustness["conditions"].items()
    }
    source_summary = {
        source: {
            "cell_count": int(count),
            "kernel_ready_by_condition": {
                condition: kernel_ready[condition][source] for condition in kernel_ready
            },
            "robustness_ready_by_condition": {
                condition: robustness_ready[condition][source] for condition in robustness_ready
            },
            "all_kernel_conditions_ready": all(
                kernel_ready[condition][source] for condition in kernel_ready
            ),
            "all_robustness_conditions_ready": all(
                robustness_ready[condition][source] for condition in robustness_ready
            ),
        }
        for source, count in config["source_types"].items()
    }
    contract = reports["external_contract"]
    minimum_train = int(contract["gates"]["minimum_recording_units_per_source_per_training_cohort"])
    minimum_validation = int(
        contract["gates"]["minimum_recording_units_per_source_per_validation_cohort"]
    )
    pseudonymous_individual_ids = one_cell_per_animal and all(paired_ordinals.values())
    training_validation_capacity = {
        source: int(count) >= minimum_train + minimum_validation
        for source, count in config["source_types"].items()
    }
    gates = {
        "all_four_T4_source_types_have_millivolt_traces": all(
            source in source_summary
            for source in contract["required_families"]["T4"]["source_types"]
        ),
        "physical_source_time_axis_available": all(
            item["response_sample_interval_milliseconds"] > 0
            for details in kernel["conditions"].values()
            for source, item in details["source_results"].items()
            if source in config["source_types"]
        ),
        "every_source_passes_both_kernel_conditions": all(
            item["all_kernel_conditions_ready"] for item in source_summary.values()
        ),
        "every_source_passes_both_robustness_conditions": all(
            item["all_robustness_conditions_ready"] for item in source_summary.values()
        ),
        "stable_pseudonymous_biological_individual_ID_available": (pseudonymous_individual_ids),
        "individual_disjoint_training_validation_split_constructible": (
            pseudonymous_individual_ids and all(training_validation_capacity.values())
        ),
        "source_type_to_MaleCNS_mapping_available": False,
        "required_split_roles_available": False,
    }
    ready = all(gates.values())
    return {
        "protocol": {
            "name": config["name"],
            "observed_on": config["observed_on"],
            "dependencies_sha256": {
                str(CONFIG): _sha256(root / CONFIG),
                str(IMPLEMENTATION): _sha256(root / IMPLEMENTATION),
                str(paper_spec["path"]): _sha256(paper_path),
                **{str(path): _sha256(root / path) for path in paths.values()},
            },
            "parameter_fit": False,
            "runtime_modified": False,
        },
        "workbook_identity_audit": {
            "path": workbook_spec["path"],
            "sha256": _sha256(workbook_path),
            "sheet_count": len(workbook.sheetnames),
            "all_sheets_visible": True,
            "custom_document_property_count": len(workbook.custom_doc_props),
            "identity_fields": identity_fields,
            "source_headers": source_headers,
            "header_semantics": "anonymous_source_type_cell_ordinal_only",
            "ON_OFF_ordinals_pair_by_source": paired_ordinals,
        },
        "paper_identity_provenance": {
            **paper_spec,
            "actual_sha256": _sha256(paper_path),
            "actual_pages": len(paper.pages),
            "embedded_attachment_names": sorted(paper.attachments),
            "one_electrophysiology_cell_per_different_animal_declared": (one_cell_per_animal),
            "source_cell_ordinals_are_pseudonymous_individual_keys": (pseudonymous_individual_ids),
        },
        "verified_Edmond_subset_identity_audit": {
            "file_count": len(verified_files),
            "files": verified_files,
            "identity_sidecars": identity_sidecars,
            "complete_dataset_manifest_rechecked": False,
            "reason": "Edmond_API_connection_timeout_during_this_audit",
        },
        "source_summary": source_summary,
        "training_validation_capacity_by_source": training_validation_capacity,
        "minimum_unique_individuals_for_training_and_validation": (
            minimum_train + minimum_validation
        ),
        "prior_group_ordering": kernel["prior_fast_delayed_grouping"],
        "transfer_gates": gates,
        "T4_individual_level_source_validation_ready": ready,
        "authorize_T4_source_dynamics_fit": ready,
        "authorize_T4_functional_precheck": ready,
        "advance_to_LPLC_mechanism_repair": False,
        "advance_to_vehicle_experiments": False,
        "blocking_requirements": [name for name, passed in gates.items() if not passed],
        "stop_reason": (
            None
            if ready
            else (
                "T4_millivolt_cells_have_pseudonymous_individual_identity_but_fail_"
                "fixed_kernel_robustness_mapping_and_split_role_gates"
            )
        ),
        "boundary": config["boundary"],
    }
